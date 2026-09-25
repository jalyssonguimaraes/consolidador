import openpyxl, json, hashlib, collections
from pathlib import Path
from decimal import Decimal

source=Path(r'C:\Users\karen\OneDrive\Desktop\1.1 - Italo Mendes\CONTROLE DE AÇÕES - ITALO (version 2).xlsx')
book=openpyxl.load_workbook(source,data_only=True)
digest=hashlib.sha256(source.read_bytes()).hexdigest()
groups=collections.OrderedDict()
def text(v):
    return format(Decimal(str(v or 0)),'f')
for row in book['Dados'].iter_rows(min_row=2,values_only=False):
    v=[c.value for c in row]
    if not v[1]: continue
    key=(v[2],v[0].strftime('%Y-%m-%d'),str(v[1]))
    if key not in groups:
        groups[key]={'id':'source-'+hashlib.sha256('|'.join(key).encode()).hexdigest()[:24],'broker':key[0],'date':key[1],'number':key[2],'kind':'variable','fees':dict.fromkeys(['settlement','exchange','brokerage','iss','irrf'],Decimal(0)),'trades':[],'source':{'file':source.name,'sha256':digest,'sheet':'Dados','rows':[]},'original':[]}
    n=groups[key]
    n['source']['rows'].append(row[0].row)
    n['trades'].append({'asset':v[6],'category':'FII' if v[5]=='Fundos Imobiliários' else 'Ação','side':'buy' if v[7]=='C' else 'sell','quantity':text(v[10]),'price':text(v[11])})
    orig={k:text(v[c]) for k,c in zip(n['fees'],range(13,18))}
    n['original'].append(orig)
    for k,amount in orig.items(): n['fees'][k]+=Decimal(amount)
for n in groups.values(): n['fees']={k:text(v.quantize(Decimal('0.01'))) for k,v in n['fees'].items()}
notes=list(groups.values())
for row in book['Tesouro Direto'].iter_rows(min_row=2):
    v=[c.value for c in row]
    if not v[2]: continue
    notes.append({'id':'treasury-'+str(v[2]),'broker':v[1],'date':v[5].strftime('%Y-%m-%d'),'number':str(v[2]),'kind':'treasury','fees':{'settlement':text(v[9]),'exchange':'0','brokerage':'0','iss':'0','irrf':'0'},'trades':[{'asset':v[3],'category':'Tesouro Direto','side':'buy' if v[4]=='Aporte' else 'sell','quantity':text(v[7]),'price':text(v[8]),'maturity':v[6].strftime('%Y-%m-%d')}],'source':{'file':source.name,'sha256':digest,'sheet':'Tesouro Direto','rows':[row[0].row]},'original':[{'b3':text(v[9]),'gross':text(v[10]),'reportedNet':text(v[11])}]})
Path('data').mkdir(exist_ok=True)
Path('data/source.json').write_text(json.dumps({'file':source.name,'sha256':digest,'notes':notes},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'notes':len(notes),'trades':sum(len(n['trades']) for n in notes),'sha256':digest}))
